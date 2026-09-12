# -*- coding: utf-8 -*-
"""CSP_考前综合笔记.md -> Word 文档（打印用）
用法：python md2docx.py
每次笔记更新后重新运行即可生成新版 docx。
"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "CSP_考前综合笔记.md"
DST = "CSP_考前综合笔记.docx"

doc = Document()
sec = doc.sections[0]
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, m, Cm(1.5))

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def shade(p, color="F2F2F2"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    p._p.get_or_add_pPr().append(shd)


def add_runs(p, text):
    """处理 **加粗** 和 `行内代码`"""
    for part in re.split(r"(\*\*.*?\*\*|`[^`]*`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x8B, 0x1A, 0x1A)
        else:
            p.add_run(part)


def code_block(code):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.4)
    pf.space_before = Pt(4)
    pf.space_after = Pt(6)
    for i, line in enumerate(code.rstrip("\n").split("\n")):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        r.font.size = Pt(8.5)
    shade(p)


quote_buf = []
table_buf = []


def flush_quote():
    global quote_buf
    if not quote_buf:
        return
    for q in quote_buf:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(2)
        add_runs(p, q)
        for r in p.runs:
            r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    quote_buf = []


def flush_table():
    global table_buf
    if not table_buf:
        return
    rows = [r for r in table_buf
            if not re.match(r"^\s*\|?[\s:\-|]+\|?\s*$", r)]
    data = []
    for r in rows:
        data.append([c.strip() for c in r.strip().strip("|").split("|")])
    ncol = max(len(c) for c in data)
    t = doc.add_table(rows=len(data), cols=ncol)
    t.style = "Table Grid"
    for ri, row in enumerate(data):
        for ci in range(ncol):
            p = t.cell(ri, ci).paragraphs[0]
            add_runs(p, row[ci] if ci < len(row) else "")
            for r in p.runs:
                r.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    table_buf = []


lines = open(SRC, encoding="utf-8").read().split("\n")
i = 0
in_code = False
code_buf = []

while i < len(lines):
    line = lines[i]
    if in_code:
        if line.strip().startswith("```"):
            code_block("\n".join(code_buf))
            code_buf = []
            in_code = False
        else:
            code_buf.append(line)
        i += 1
        continue
    if line.strip().startswith("```"):
        in_code = True
        i += 1
        continue
    if line.startswith("#"):
        flush_quote()
        flush_table()
        m = re.match(r"^(#+)\s*(.*)", line)
        level = len(m.group(1))
        h = doc.add_heading(m.group(2), level=min(level, 3))
        for r in h.runs:
            r.font.color.rgb = RGBColor(0, 0, 0)
            r.font.name = "Calibri"
            r.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            r.font.size = Pt({1: 15, 2: 12.5, 3: 11}.get(level, 11))
        i += 1
        continue
    if line.startswith(">"):
        flush_table()
        quote_buf.append(line.lstrip("> "))
        i += 1
        continue
    if line.strip().startswith("|"):
        flush_quote()
        table_buf.append(line)
        i += 1
        continue
    if line.strip() in ("", "---"):
        flush_quote()
        flush_table()
        i += 1
        continue
    flush_quote()
    flush_table()
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    if re.match(r"^\s*[-*]\s", line):
        p.paragraph_format.left_indent = Cm(0.5)
        add_runs(p, "\u2022 " + re.sub(r"^\s*[-*]\s+", "", line))
    elif re.match(r"^\s*\d+\.\s", line):
        p.paragraph_format.left_indent = Cm(0.5)
        add_runs(p, line.strip())
    else:
        add_runs(p, line)
    i += 1

flush_quote()
flush_table()
doc.save(DST)
print("已生成:", DST)
